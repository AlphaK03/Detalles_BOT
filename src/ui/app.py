"""app.py — Ventana principal del bot (CHC CONSULTORES — Detalles BOT).

Pantalla compacta con selección, progreso, consola y acciones.
Configuración separada; administrador local de tablas en ventana aparte.
"""
import os
import sys
import threading
import logging
import queue
from tkinter import ttk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
import tkinter as tk
from openpyxl import load_workbook

_SRC_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _SRC_DIR not in sys.path:
    sys.path.insert(0, _SRC_DIR)

from utils import (
    load_machine_config, save_machine_config,
    load_session, save_session,
    compute_output_paths, get_catalog_folder, dataset_key, setup_logging,
)
from core.playwright_processor import PlaywrightDUAProcessor
from core.captcha import load_captcha_lookup
from core.job_store import JobStore, RunLock
from core.catalog import TableCatalog
from flows.detalles_flow import DetallesFlow
from ui.theme import apply_theme
from ui.database_manager import DatabaseWindow
from ui.import_dialog import ImportNameDialog


# ── Logging handler ───────────────────────────────────────────────────────────

class _GuiLogHandler(logging.Handler):
    _FMT = logging.Formatter('%(asctime)s  %(message)s',
                              datefmt='%H:%M:%S')

    def __init__(self, events):
        super().__init__()
        self.events = events
        self.setFormatter(self._FMT)

    def emit(self, record: logging.LogRecord) -> None:
        if getattr(record, 'technical_only', False):
            return
        msg = self.format(record)
        # Permite tags personalizados via extra={'tag': 'SUCCESS'}
        tag = getattr(record, 'tag', record.levelname)

        self.events.put(('log', (msg, tag)))


# ── App ───────────────────────────────────────────────────────────────────────

class App(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("CHC CONSULTORES — Detalles BOT")
        self.geometry("760x610")
        self.minsize(720, 570)

        self.machine = load_machine_config()
        self.session = load_session()
        self.catalog = TableCatalog(get_catalog_folder())
        self.catalog.discover_legacy()
        self._active_table = None
        self._importing = False

        self._captcha_lookup = None
        self._events = queue.SimpleQueue()
        self._running = False
        self._closing = False
        self._stop_event = threading.Event()
        self._stats = {}

        self._build_ui()
        self._setup_logging()
        try:
            self._captcha_lookup = load_captcha_lookup()
            logging.info('Aplicación lista para consultar.')
        except Exception as exc:
            messagebox.showerror('Catálogo CAPTCHA', str(exc))
        self._restore_session()
        self.after(100, self._poll_events)

        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ═════════════════════════════════════════════════════════════════════════
    # Construcción de la interfaz
    # ═════════════════════════════════════════════════════════════════════════

    def _build_ui(self):
        self._dark_var = tk.BooleanVar(value=self.machine.get('dark_mode', True))
        self.colors = apply_theme(self, self._dark_var.get())
        self._config_widgets = []
        self._config_dirty = False
        header = ttk.Frame(self, padding=(12, 8, 12, 5))
        header.pack(fill='x')
        ttk.Label(header, text="Detalles BOT", font=('Segoe UI', 13, 'bold')).pack(side='left')
        ttk.Label(header, text="CHC CONSULTORES", style='Muted.TLabel').pack(side='right')
        self._notebook = ttk.Notebook(self)
        self._tab_proc = ttk.Frame(self._notebook, padding=8)
        self._tab_cfg = ttk.Frame(self._notebook, padding=8)
        self._notebook.add(self._tab_proc, text="Procesamiento")
        self._notebook.add(self._tab_cfg, text="Configuración")
        self._path_vars = {'output_base_path': tk.StringVar(value=self.machine.get('output_base_path', ''))}
        self._build_tab_procesamiento(self._tab_proc)
        self._build_tab_configuracion(self._tab_cfg)
        self._build_action_bar()
        self._notebook.pack(fill='both', expand=True, padx=10)
        for var in (*self._param_vars.values(), self._headless_var):
            var.trace_add('write', self._mark_config_dirty)
        self._dark_var.trace_add('write', self._appearance_changed)

    def _build_tab_procesamiento(self, parent):
        fields = ttk.LabelFrame(parent, text=" Importar o abrir una tabla guardada ", padding=(8, 6))
        fields.pack(fill='x', pady=(0, 7))
        fields.columnconfigure(1, weight=1)
        self._file_var = tk.StringVar()
        self._file_display = tk.StringVar(value="Selecciona un archivo Excel")
        ttk.Label(fields, text="Archivo:", style='Card.TLabel').grid(row=0, column=0, sticky='w', padx=(0, 8))
        ttk.Entry(fields, textvariable=self._file_var, state='readonly').grid(row=0, column=1, sticky='ew')
        self._browse_btn = ttk.Button(fields, text="Elegir Excel…", command=self._select_file, width=15)
        self._browse_btn.grid(row=0, column=2, padx=(8, 0), pady=2)
        ttk.Label(fields, text="Hoja:", style='Card.TLabel').grid(row=1, column=0, sticky='w')
        self._sheet_var = tk.StringVar(value=self.session.get('sheet_name', ''))
        self._sheet_combo = ttk.Combobox(fields, textvariable=self._sheet_var, state='disabled', width=26)
        self._sheet_combo.grid(row=1, column=1, sticky='w', pady=4)
        self._sheet_combo.bind('<<ComboboxSelected>>', self._on_sheet_selected)
        self._source_hint = ttk.Label(fields, text="", style='Hint.TLabel')
        self._source_hint.grid(row=1, column=2, sticky='w', padx=(8, 0))
        ttk.Label(fields, text="Guardar en:", style='Card.TLabel').grid(row=2, column=0, sticky='w')
        ttk.Entry(fields, textvariable=self._path_vars['output_base_path'], state='readonly').grid(
            row=2, column=1, sticky='ew')
        self._output_btn = ttk.Button(fields, text="Elegir carpeta…", width=15,
                                      command=lambda: self._browse_folder('output_base_path'))
        self._output_btn.grid(row=2, column=2, padx=(8, 0), pady=2)
        ttk.Label(fields, text="Tabla:", style='Card.TLabel').grid(row=3, column=0, sticky='w')
        self._table_name = tk.StringVar()
        self._name_entry = ttk.Entry(fields, textvariable=self._table_name)
        self._name_entry.grid(row=3, column=1, sticky='ew', pady=4)
        self._import_btn = ttk.Button(fields, text="Guardar tabla", command=self._import_table, width=15)
        self._import_btn.grid(row=3, column=2, padx=(8, 0))
        self._file_var.trace_add('write', self._update_source_label)

        progress = ttk.LabelFrame(parent, text=" Progreso ", padding=(8, 6))
        progress.pack(fill='x', pady=(0, 7))
        top = ttk.Frame(progress, style='Card.TFrame')
        top.pack(fill='x')
        self._counter_label = ttk.Label(top, text="0 / —", style='Card.TLabel',
                                        font=('Segoe UI', 17, 'bold'))
        self._counter_label.pack(side='left')
        self._pct_label = ttk.Label(top, text="0.0%", style='Card.TLabel',
                                    foreground=self.colors['ACCENT'], font=('Segoe UI', 11, 'bold'))
        self._pct_label.pack(side='right')
        self._progress_var = tk.DoubleVar()
        ttk.Progressbar(progress, variable=self._progress_var, maximum=100).pack(fill='x', pady=(4, 6))
        metrics = ttk.Frame(progress, style='Card.TFrame')
        metrics.pack(fill='x')
        self._metric_labels = {}
        for column, (key, title, color) in enumerate([
            ('succeeded', 'Guardados:', self.colors['SUCCESS']), ('remaining', 'Pendientes:', self.colors['INK']),
            ('running', 'En curso:', self.colors['ACCENT']), ('invalid', 'Revisar:', self.colors['WARNING']),
        ]):
            metrics.columnconfigure(column, weight=1)
            cell = ttk.Frame(metrics, style='Card.TFrame')
            cell.grid(row=0, column=column, sticky='w')
            ttk.Label(cell, text=title, style='Hint.TLabel').pack(side='left')
            value = ttk.Label(cell, text="—", style='Card.TLabel', foreground=color,
                              font=('Segoe UI', 10, 'bold'))
            value.pack(side='left', padx=(4, 0))
            self._metric_labels[key] = value
        self._status_label = ttk.Label(progress, text="Listo para empezar", style='Card.TLabel',
                                       font=('Segoe UI', 9, 'bold'), foreground=self.colors['MUTED'])
        self._status_label.pack(anchor='w', pady=(5, 0))
        self._stats_label = ttk.Label(progress, text="Los resultados se guardan automáticamente.",
                                      style='Hint.TLabel')
        self._stats_label.pack(anchor='w', pady=(2, 0))
        self._build_activity(parent)

    def _build_tab_configuracion(self, parent):
        panel = ttk.LabelFrame(parent, text=" Parámetros de consulta ", padding=12)
        panel.pack(fill='x')
        panel.columnconfigure(2, weight=1)
        self._param_vars = {}
        for row, (key, label, hint, low, high, default) in enumerate([
            ('cant_hilos', 'Consultas simultáneas', 'Empieza con 1.', 1, 20, 1),
            ('timeout', 'Tiempo de espera (s)', 'Máximo por operación.', 5, 120, 30),
            ('max_retries', 'Rondas de intentos', 'Cada ronda prueba los pendientes.', 1, 10, 3),
        ]):
            ttk.Label(panel, text=label, style='Card.TLabel').grid(row=row, column=0, sticky='w', pady=7)
            var = tk.IntVar(value=int(self.machine.get(key, default)))
            self._param_vars[key] = var
            spin = ttk.Spinbox(panel, from_=low, to=high, textvariable=var, width=6, state='readonly')
            spin.grid(row=row, column=1, padx=14)
            ttk.Label(panel, text=hint, style='Hint.TLabel').grid(row=row, column=2, sticky='w')
            self._config_widgets.append((spin, 'readonly'))
        self._headless_var = tk.BooleanVar(value=bool(self.machine.get('headless', True)))
        toggle = ttk.Checkbutton(panel, text="Mantener el navegador oculto", variable=self._headless_var)
        toggle.grid(row=3, column=0, columnspan=3, sticky='w', pady=(12, 2))
        self._config_widgets.append((toggle, 'normal'))
        appearance = ttk.LabelFrame(parent, text=" Apariencia ", padding=12)
        appearance.pack(fill='x', pady=(10, 0))
        ttk.Checkbutton(appearance, text="Modo oscuro", variable=self._dark_var).pack(anchor='w')
        ttk.Label(appearance, text="Se aplica al instante y se recuerda al guardar la configuración.",
                  style='Hint.TLabel').pack(anchor='w', pady=(5, 0))
        row = ttk.Frame(parent)
        row.pack(fill='x', pady=10)
        self._save_btn = ttk.Button(row, text="Guardar configuración", command=self._save_config,
                                    style='Primary.TButton')
        self._save_btn.pack(side='right')
        self._cfg_saved_label = ttk.Label(row, text="Configuración guardada", style='Muted.TLabel')
        self._cfg_saved_label.pack(side='left')
        ttk.Label(parent, text="Los ajustes se aplican al iniciar. Los resultados guardados se conservan.",
                  style='Muted.TLabel', wraplength=630).pack(anchor='w', pady=4)

    def _build_activity(self, parent):
        panel = ttk.LabelFrame(parent, text=" Actividad ", padding=5)
        panel.pack(fill='both', expand=True)
        self._log_text = ScrolledText(panel, height=6, state='disabled', font=('Consolas', 9),
                                     wrap='word', background='#121B28', foreground='#F2B269',
                                     insertbackground='white', relief='sunken', borderwidth=1, padx=6, pady=5)
        self._log_text.pack(fill='both', expand=True)
        for tag, color in [('DEBUG', '#A8B7C9'), ('INFO', '#F2B269'), ('WARNING', '#FFD479'),
                           ('ERROR', '#FF8E88'), ('CRITICAL', '#FF8E88'), ('SUCCESS', '#84D8A3')]:
            self._log_text.tag_config(tag, foreground=color)

    def _build_action_bar(self):
        footer = ttk.Frame(self, padding=(12, 7, 12, 8))
        footer.pack(side='bottom', fill='x')
        self._action_hint = ttk.Label(footer, text="Selecciona un Excel para empezar.",
                                      style='Muted.TLabel', wraplength=690)
        self._action_hint.pack(anchor='w', pady=(0, 6))
        buttons = ttk.Frame(footer)
        buttons.pack(fill='x')
        self._folder_btn = ttk.Button(buttons, text="Resultados", command=self._open_output_folder)
        self._folder_btn.pack(side='left')
        self._pending_btn = ttk.Button(buttons, text="Base de datos", command=self._open_pending)
        self._pending_btn.pack(side='left', padx=5)
        self._refresh_btn = ttk.Button(buttons, text="Actualizar", command=self._refresh_counter)
        self._refresh_btn.pack(side='left')
        self._clear_btn = ttk.Button(buttons, text="Limpiar vista", command=self._clear_log)
        self._clear_btn.pack(side='left', padx=5)
        self._start_btn = ttk.Button(buttons, text="Iniciar consultas", style='Primary.TButton',
                                     command=self._start_processing)
        self._start_btn.pack(side='right')
        self._stop_btn = ttk.Button(buttons, text="Detener", style='Stop.TButton',
                                    command=self._stop_processing, state='disabled')
        self._stop_btn.pack(side='right', padx=6)

    def _clear_log(self):
        self._log_text.config(state='normal')
        self._log_text.delete('1.0', 'end')
        self._log_text.config(state='disabled')

    def _mark_config_dirty(self, *args):
        self._config_dirty = True
        self._cfg_saved_label.config(text="Cambios sin guardar", foreground=self.colors['WARNING'])

    def _appearance_changed(self, *args):
        self.colors = apply_theme(self, self._dark_var.get())
        self._pct_label.configure(foreground=self.colors['ACCENT'])
        for key, color in [('succeeded', 'SUCCESS'), ('remaining', 'INK'), ('running', 'ACCENT'), ('invalid', 'WARNING')]:
            self._metric_labels[key].configure(foreground=self.colors[color])
        self._status_label.configure(foreground=self.colors['ACCENT'])
        self._mark_config_dirty()
        for child in self.winfo_children():
            if isinstance(child, tk.Toplevel):
                child.configure(background=self.colors['BG'])

    def _update_source_label(self, *args):
        path = self._file_var.get()
        name = os.path.basename(path) if path else "Ningún archivo seleccionado"
        self._file_display.set(name if len(name) < 53 else name[:24] + '…' + name[-24:])
        self._source_hint.config(text="Tabla guardada" if self._active_table else "Columna: NroDUA")

    def _update_actions(self):
        path, sheet = self._file_var.get(), self._sheet_var.get().strip()
        base = self.machine.get('output_base_path', '')
        total, done = self._stats.get('total', 0), self._stats.get('succeeded', 0)
        if self._running:
            self._start_btn.config(text="Importando…" if self._importing else "Consultando…", state='disabled')
            self._action_hint.config(text="Guardando la lista; espera un momento." if self._importing else
                                     "Puedes detenerte cuando quieras. Los pendientes se conservan.")
            return
        if not self._active_table:
            hint = "Elige Excel y hoja, revisa el nombre y guarda la tabla. O abre una desde Base de datos."
        elif not base:
            hint = "Elige dónde guardar los resultados."
        elif total and done == total:
            hint = "Todos los DUAs tienen resultado. Puedes volver a generar los archivos."
        elif total:
            hint = "Retoma los pendientes; los resultados guardados se conservan."
        else:
            hint = "Tabla guardada. Ya no necesitas el archivo Excel para continuar."
        self._action_hint.config(text=hint)
        text = "Actualizar resultados" if total and done == total else (
            "Retomar pendientes" if total else "Iniciar consultas")
        ready = bool(self._active_table and base and self._captcha_lookup is not None)
        self._start_btn.config(text=text, state='normal' if ready else 'disabled')
        self._folder_btn.config(state='normal' if base and self._active_table else 'disabled')
        self._pending_btn.config(state='normal')
        self._import_btn.config(state='normal' if path and sheet and not self._active_table else 'disabled')
        self._name_entry.config(state='readonly' if self._active_table else 'normal')

    def _open_pending(self):
        DatabaseWindow(self)

    def _setup_logging(self):
        setup_logging()
        handler = _GuiLogHandler(self._events)
        handler.setLevel(logging.INFO)
        logging.getLogger().addHandler(handler)
        self._log_handler = handler

    # ═════════════════════════════════════════════════════════════════════════
    # Sesión
    # ═════════════════════════════════════════════════════════════════════════

    def _restore_session(self):
        path = self.session.get('input_file_path', '')
        sheet = self.session.get('sheet_name', '')
        table_id = self.session.get('active_table_id') or (dataset_key(path, sheet) if path and sheet else '')
        if table_id:
            try:
                self._activate_table(self.catalog.get(table_id))
                return
            except (ValueError, OSError):
                pass
        self._file_var.set(path)
        if path and os.path.exists(path):
            self._populate_sheets(path, sheet)
        self._refresh_counter()

    def _activate_table(self, record):
        if self._running:
            return
        self._active_table = self.catalog.get(record['id'])
        self._file_var.set('Lista guardada · Excel original ya no necesario')
        self._sheet_var.set(record['source_sheet'])
        self._sheet_combo.config(values=[record['source_sheet']], state='disabled')
        self._table_name.set(record['name'])
        self._save_selection()
        self._refresh_counter()

    def _catalog_changed(self):
        if self._active_table:
            try:
                self._activate_table(self.catalog.get(self._active_table['id']))
            except (ValueError, OSError):
                self._active_table = None
                self._file_var.set('')
                self._sheet_var.set('')
                self._table_name.set('')
                self._sheet_combo.config(values=[], state='disabled')
                self._save_selection()
                self._refresh_counter()

    def _output_paths(self):
        record = self._active_table
        return compute_output_paths(self.machine['output_base_path'], record['output_name'],
                                    table_id=record['id'])

    def _import_table(self):
        if self._running or self._active_table:
            return
        source, sheet, name = self._file_var.get(), self._sheet_var.get(), self._table_name.get()
        choice = ImportNameDialog(self, name).result
        if choice is None:
            return
        name, expected_target = choice
        self._table_name.set(name)
        base = self.machine.get('output_base_path', '')
        legacy_csv = os.path.join(base, sheet, f'MM20YY_Detalles_DUAs__procesados_{sheet}.csv') if base else None
        self._importing = True
        self._set_controls_running(True)
        self._status_label.config(text="Guardando la tabla…", foreground=self.colors['ACCENT'])

        def run():
            try:
                record = self.catalog.import_excel(source, sheet, name, legacy_csv,
                                                   list(DetallesFlow().extract_data('').columns),
                                                   expected_target=expected_target)
                self._events.put(('imported', (record, None)))
            except Exception as exc:
                self._events.put(('imported', (None, str(exc))))

        self._worker_thread = threading.Thread(target=run, daemon=True)
        self._worker_thread.start()

    def _refresh_counter(self):
        if self._running:
            return
        database = self.catalog.database(self._active_table) if self._active_table else ''
        if database and os.path.exists(database):
            store = JobStore(database)
            try:
                self._show_summary(store.summary())
            finally:
                store.close()
        else:
            self._stats = {}
            self._counter_label.config(text="0 / —")
            self._pct_label.config(text="0.0%")
            self._progress_var.set(0)
            for label in self._metric_labels.values():
                label.config(text='—')
            self._stats_label.config(text="Los resultados se guardan automáticamente.")
            self._status_label.config(text="Listo para empezar", foreground=self.colors['MUTED'])
        self._update_source_label()
        self._update_actions()

    def _select_file(self):
        path = filedialog.askopenfilename(
            title="Seleccionar archivo de DUAs",
            filetypes=[("Excel", "*.xlsx"), ("Todos los archivos", "*.*")])
        if not path:
            return
        self._active_table = None
        self._file_var.set(path)
        self._populate_sheets(path, self.session.get('sheet_name', ''))
        self._save_selection()
        self._refresh_counter()

    def _save_selection(self):
        self.session['active_table_id'] = self._active_table['id'] if self._active_table else ''
        self.session['input_file_path'] = self._active_table['source_path'] if self._active_table else self._file_var.get()
        self.session['sheet_name'] = self._active_table['source_sheet'] if self._active_table else self._sheet_var.get().strip()
        save_session(self.session)

    def _on_sheet_selected(self, event=None):
        self._table_name.set(self._sheet_var.get())
        self._save_selection()
        self._refresh_counter()

    def _populate_sheets(self, path, preferred=''):
        self._sheet_var.set('')
        self._sheet_combo.config(values=[], state='disabled')
        try:
            workbook = load_workbook(path, read_only=True)
            try:
                sheets = workbook.sheetnames
            finally:
                workbook.close()
        except Exception as exc:
            messagebox.showerror("Error", f"No se pudo abrir el archivo:\n{exc}")
            return
        self._sheet_combo.config(values=sheets, state='readonly')
        self._sheet_var.set(preferred if preferred in sheets else (sheets[0] if sheets else ''))
        self._table_name.set(self._sheet_var.get())

    def _open_output_folder(self):
        base = self.machine.get('output_base_path', '')
        if not base:
            messagebox.showwarning("Resultados", "Configure una carpeta de salida.")
            return
        if not self._active_table:
            return
        paths = self._output_paths()
        try:
            os.startfile(os.path.dirname(paths['output_csv']))
        except OSError as exc:
            messagebox.showerror("No se pudo abrir la carpeta", str(exc))

    def _browse_folder(self, key):
        if self._running:
            return
        folder = filedialog.askdirectory(title="Seleccionar carpeta")
        if folder:
            updated = {**self.machine, key: folder}
            try:
                save_machine_config(updated)
            except OSError as exc:
                messagebox.showerror("No se pudo guardar la carpeta", str(exc))
                return
            self.machine = updated
            self._path_vars[key].set(folder)
            self._update_actions()

    def _save_config(self):
        if self._running:
            return False
        updated = dict(self.machine)
        for key, var in self._param_vars.items():
            try:
                value = int(var.get())
                low, high = {'cant_hilos': (1, 20), 'timeout': (5, 120), 'max_retries': (1, 10)}[key]
                if not low <= value <= high:
                    raise ValueError()
                updated[key] = value
            except (ValueError, tk.TclError):
                messagebox.showerror("Error", f"Valor inválido para '{key}'.")
                return False
        for key, var in self._path_vars.items():
            updated[key] = var.get().strip()
        updated['headless'] = bool(self._headless_var.get())
        updated['dark_mode'] = bool(self._dark_var.get())
        # Claves heredadas que ya no usa el motor activo.
        for key in ('names_file_path', 'edge_driver_path', 'screenshot_path'):
            updated.pop(key, None)
        try:
            save_machine_config(updated)
        except OSError as exc:
            messagebox.showerror("No se pudieron guardar los ajustes", str(exc))
            return False
        self.machine = updated
        self._config_dirty = False
        self._cfg_saved_label.config(text="Ajustes guardados", foreground=self.colors['SUCCESS'])
        self._update_actions()
        logging.info("Configuración guardada.")
        return True

    def _start_processing(self):
        if self._running:
            return
        if self._config_dirty and not self._save_config():
            self._notebook.select(self._tab_cfg)
            return
        if self._captcha_lookup is None:
            messagebox.showerror("Error", "No se cargó el catálogo interno de CAPTCHAs.")
            return
        base = self.machine.get('output_base_path', '')
        if not self._active_table:
            messagebox.showwarning("Atención", "Guarda la tabla o abre una desde Base de datos.")
            return
        if not base:
            messagebox.showwarning("Atención", "Configure la carpeta de salida.")
            return
        try:
            self._save_selection()
            paths = self._output_paths()
        except OSError as exc:
            messagebox.showerror("No se pudo preparar la consulta", str(exc))
            return
        self._stop_event = threading.Event()
        self._set_controls_running(True)
        self._notebook.select(self._tab_proc)
        self._status_label.config(text="Preparando tus DUAs…", foreground=self.colors['ACCENT'])
        self._worker_thread = threading.Thread(
            target=self._run, args=(self._active_table['id'], paths, dict(self.machine)), daemon=True)
        self._worker_thread.start()

    def _stop_processing(self):
        self._stop_event.set()
        self._status_label.config(text="Deteniendo y guardando…", foreground=self.colors['WARNING'])
        self._action_hint.config(text="Espera un momento mientras se guardan los pendientes.")
        self._stop_btn.config(state='disabled')

    def _set_controls_running(self, running):
        self._running = running
        for widget in (self._start_btn, self._browse_btn, self._save_btn, self._refresh_btn,
                       self._output_btn, self._import_btn, self._name_entry):
            widget.config(state='disabled' if running else 'normal')
        for widget, idle_state in self._config_widgets:
            widget.config(state='disabled' if running else idle_state)
        self._stop_btn.config(state='normal' if running and not self._importing else 'disabled')
        self._sheet_combo.config(state='disabled' if running or self._active_table or not self._sheet_var.get() else 'readonly')
        self._update_actions()

    def _run(self, table_id, paths, machine):
        error = None
        summary = None
        try:
            database = self.catalog.database(self.catalog.get(table_id))
            with RunLock(database):
                self.catalog.get(table_id)
                store = JobStore(database)
                try:
                    flow = DetallesFlow()
                    logging.info('Tabla lista: %s DUAs; %s ya guardados.',
                                 store.summary()['total'], store.summary()['succeeded'])
                    processor = PlaywrightDUAProcessor(
                        machine, paths, flow, self._captcha_lookup, self._stop_event)
                    processor.run(store, lambda value: self._events.put(('summary', value)))
                finally:
                    summary = store.summary()
                    store.close()
        except Exception as exc:
            error = ('No se pudieron leer o guardar los archivos. Revisa la carpeta de resultados.'
                     if isinstance(exc, OSError) else
                     'La consulta se interrumpió. Los pendientes y resultados guardados se conservan.')
            logging.error(error)
            logging.exception('Ejecución detenida por error: %s', exc, extra={'technical_only': True})
        finally:
            self._events.put(('finished', (summary, error)))

    def _poll_events(self):
        for _ in range(200):
            try:
                kind, value = self._events.get_nowait()
            except queue.Empty:
                break
            if kind == 'log':
                msg, tag = value
                self._log_text.config(state='normal')
                self._log_text.insert('end', msg + '\n', tag)
                if int(self._log_text.index('end-1c').split('.')[0]) > 2000:
                    self._log_text.delete('1.0', '201.0')
                self._log_text.see('end')
                self._log_text.config(state='disabled')
            elif kind == 'summary':
                self._show_summary(value)
            elif kind == 'finished':
                self._on_finished(*value)
            elif kind == 'imported':
                record, error = value
                self._importing = False
                self._set_controls_running(False)
                if record:
                    self._activate_table(record)
                    report = record['import_summary']
                    text = (f'Tabla «{record["name"]}»: {report["added"]:,} DUAs nuevos, '
                            f'{report["existing"]:,} ya registrados y {report["invalid"]:,} entradas para revisar.')
                    logging.info(text)
                    self._action_hint.config(text=text)
                    if not self._closing:
                        messagebox.showinfo('Tabla actualizada' if record['appended'] else 'Tabla guardada',
                                            text + '\n\nYa puedes continuar sin el Excel original.', parent=self)
                elif not self._closing:
                    self._status_label.config(text="No se pudo guardar la tabla", foreground=self.colors['ERROR'])
                    messagebox.showerror('Importar tabla', error)
        if self._closing and not self._running:
            logging.getLogger().removeHandler(self._log_handler)
            self.destroy()
            return
        self.after(100, self._poll_events)

    def _show_summary(self, summary):
        self._stats = summary
        total, done = summary['total'], summary['succeeded']
        pct = done / total * 100 if total else 0
        self._counter_label.config(text=f"{done:,} / {total:,}")
        self._pct_label.config(text=f"{pct:.1f}%")
        self._progress_var.set(pct)
        values = {**summary, 'remaining': total - done}
        for key, label in self._metric_labels.items():
            label.config(text=f"{values[key]:,}")
        if summary['invalid']:
            note = f"{summary['invalid']:,} entradas no son DUAs válidos. Revísalas en Base de datos."
        else:
            note = (f"{summary['retry']:,} pendientes de reintento automático al terminar la lista."
                    if self._running and summary['retry'] else "Los resultados se guardan automáticamente.")
        self._stats_label.config(text=note)
        stopped = self._running and self._stop_event.is_set()
        text = 'Deteniendo y guardando…' if stopped else (
            'Consultando el portal…' if self._running else (
                'Todo guardado' if total and total == done else 'Pendientes listos para retomar'))
        self._status_label.config(text=text, foreground=self.colors['ACCENT'] if self._running else (
            self.colors['WARNING'] if total > done else self.colors['SUCCESS']))
        if not stopped:
            self._update_actions()

    def _on_finished(self, summary, error):
        self._set_controls_running(False)
        if summary is not None:
            self._show_summary(summary)
        if error:
            title, color = 'No se pudo terminar la consulta', self.colors['ERROR']
        elif self._stop_event.is_set():
            title, color = 'Detenido; pendientes guardados', self.colors['WARNING']
        elif summary and summary['total'] == summary['succeeded']:
            title, color = 'Completado', self.colors['SUCCESS']
        else:
            title, color = 'Finalizado con pendientes', self.colors['WARNING']
        self._status_label.config(text=title, foreground=color)
        if self._closing:
            return
        if error:
            self._action_hint.config(text="Revisa Actividad para ver el error. Los resultados guardados se conservan.")
            messagebox.showerror(title, error)

    def _on_close(self):
        self._closing = True
        if self._running:
            self._stop_processing()
        # _poll_events espera cierre de workers y exportación antes de destruir Tk.
