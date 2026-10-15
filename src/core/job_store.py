"""Cola durable. SQLite es la fuente de verdad; CSV/JSON son exportaciones."""
import csv
import hashlib
import json
import os
import sqlite3
import tempfile
import threading
import uuid
from contextlib import contextmanager
from pathlib import Path

from core.identifiers import normalize_dua


def active_revisions(db):
    row = db.execute("SELECT value FROM metadata WHERE key='active_revisions'").fetchone()
    if row:
        return json.loads(row[0])
    row = db.execute("SELECT value FROM metadata WHERE key='revision'").fetchone()
    return [json.loads(row[0]) if row else '']


class RunLock:
    """Una ejecución por lote; el SO libera el bloqueo incluso tras un crash."""
    def __init__(self, database):
        self.path = Path(str(database) + '.lock')
        self.stream = None

    def __enter__(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.stream = self.path.open('a+b')
        if self.path.stat().st_size == 0:
            self.stream.write(b'0')
            self.stream.flush()
        self.stream.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(self.stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.stream.close()
            raise RuntimeError('Este lote ya está siendo procesado por otra instancia.') from exc
        return self

    def __exit__(self, *args):
        self.stream.seek(0)
        if os.name == 'nt':
            import msvcrt
            msvcrt.locking(self.stream.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            fcntl.flock(self.stream, fcntl.LOCK_UN)
        self.stream.close()


@contextmanager
def atomic_output(path, encoding='utf-8', newline=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(prefix=path.name + '.', suffix='.tmp', dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding=encoding, newline=newline) as stream:
            yield stream
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


class JobStore:
    def __init__(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.path = str(path)
        self.lock = threading.RLock()
        self.db = sqlite3.connect(self.path, timeout=30, check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.execute('PRAGMA synchronous=FULL')
        self.db.executescript('''
            CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS jobs (
                id TEXT PRIMARY KEY, dua TEXT, status TEXT NOT NULL
                    CHECK(status IN ('pending','running','retry','succeeded','invalid')),
                attempts INTEGER NOT NULL DEFAULT 0, run_attempts INTEGER NOT NULL DEFAULT 0,
                last_error TEXT NOT NULL DEFAULT '', error_kind TEXT NOT NULL DEFAULT '',
                evidence TEXT NOT NULL DEFAULT '', result TEXT,
                CHECK ((status = 'succeeded') = (result IS NOT NULL))
            );
            CREATE INDEX IF NOT EXISTS jobs_queue ON jobs(status, run_attempts);
            CREATE UNIQUE INDEX IF NOT EXISTS jobs_unique_dua ON jobs(dua) WHERE dua IS NOT NULL;
            CREATE INDEX IF NOT EXISTS jobs_ready ON jobs(run_attempts)
                WHERE status IN ('pending','retry');
            CREATE TABLE IF NOT EXISTS source_rows (
                revision TEXT NOT NULL, row_number INTEGER NOT NULL, raw TEXT NOT NULL,
                job_id TEXT NOT NULL REFERENCES jobs(id), PRIMARY KEY(revision,row_number)
            );
            CREATE TABLE IF NOT EXISTS attempts (
                id INTEGER PRIMARY KEY, job_id TEXT NOT NULL REFERENCES jobs(id),
                started TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, finished TEXT,
                outcome TEXT NOT NULL DEFAULT 'running', error TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS source_job ON source_rows(job_id,revision);
            CREATE TABLE IF NOT EXISTS retained_results (
                dua TEXT PRIMARY KEY, result TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS imports (
                revision TEXT PRIMARY KEY, source TEXT NOT NULL, sheet TEXT NOT NULL,
                rows INTEGER NOT NULL, added INTEGER NOT NULL, existing INTEGER NOT NULL,
                invalid INTEGER NOT NULL, created TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
        ''')

    def close(self):
        self.db.close()

    @contextmanager
    def transaction(self):
        with self.lock:
            self.db.execute('BEGIN IMMEDIATE')
            try:
                yield self.db
                self.db.commit()
            except BaseException:
                self.db.rollback()
                raise

    def meta(self, key, default=None):
        with self.lock:
            row = self.db.execute('SELECT value FROM metadata WHERE key=?', (key,)).fetchone()
        return json.loads(row[0]) if row else default

    @staticmethod
    def _set_meta(db, key, value):
        db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', (key, json.dumps(value)))

    def import_values(self, values, revision='manual', signature=None, append=False):
        """Una transacción: o entra toda la hoja o no cambia nada. Conserva filas inválidas."""
        count = 0
        added = existing = invalid = 0
        with self.transaction() as db:
            revisions = active_revisions(db) if append else []
            for count, value in enumerate(values, start=1):
                raw = '' if value is None else str(value).strip()
                try:
                    dua = normalize_dua(raw)
                    job_id, status, error = dua, 'pending', ''
                except ValueError as exc:
                    invalid += 1
                    dua, status, error = None, 'invalid', str(exc)
                    job_id = 'invalid:' + hashlib.sha256(f'{count}:{raw}'.encode()).hexdigest()
                if append and dua and db.execute('SELECT 1 FROM retained_results WHERE dua=?', (dua,)).fetchone():
                    existing += 1
                    continue
                inserted = db.execute('''INSERT OR IGNORE INTO jobs(id,dua,status,last_error,error_kind)
                              VALUES (?,?,?,?,?)''',
                           (job_id, dua, status, error, 'input' if error else '')).rowcount
                if dua:
                    added += inserted
                    existing += 1 - inserted
                db.execute('INSERT OR REPLACE INTO source_rows VALUES (?,?,?,?)',
                           (revision, count + 1, raw, job_id))
            if count == 0:
                raise ValueError('La hoja no contiene filas de DUAs.')
            self._set_meta(db, 'source_rows', count + (self.meta('source_rows', 0) if append else 0))
            self._set_meta(db, 'revision', revision)
            self._set_meta(db, 'active_revisions', list(dict.fromkeys([*revisions, revision])))
            report = dict(rows=count, added=added, existing=existing, invalid=invalid)
            self._set_meta(db, 'last_import', report)
            if signature is not None:
                self._set_meta(db, 'source_signature', signature)
                db.execute('''INSERT OR REPLACE INTO imports(revision,source,sheet,rows,added,existing,invalid)
                              VALUES (?,?,?,?,?,?,?)''', (revision, signature[0], signature[1], count, added, existing, invalid))
        return self.summary()

    def import_excel(self, path, sheet, append=False):
        """Excel solo en la frontera. Reanudar usa SQLite si el archivo no cambió."""
        from openpyxl import load_workbook
        path = Path(path)
        if not path.exists():
            if not append and self.meta('revision') is not None:
                return self.summary()
            raise FileNotFoundError(path)
        stat = path.stat()
        signature = [str(path.resolve()), sheet, stat.st_size, stat.st_mtime_ns]
        if not append and self.meta('source_signature') == signature:
            return self.summary()
        revision = uuid.uuid4().hex if append else hashlib.sha256(json.dumps(signature).encode()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook[sheet].iter_rows(values_only=True)
            header = next(rows, ())
            if header.count('NroDUA') != 1:
                raise ValueError('La hoja debe contener exactamente una columna NroDUA.')
            column = header.index('NroDUA')
            return self.import_values((row[column] for row in rows), revision, signature, append=append)
        finally:
            workbook.close()

    def migrate_csv(self, path, fields):
        """Solo filas completas y correspondientes al lote; un log nunca prueba éxito."""
        if self.meta('legacy_migrated') or not path or not Path(path).exists():
            return 0
        imported = 0
        with open(path, encoding='utf-8-sig', newline='') as stream, self.transaction() as db:
            reader = csv.DictReader(stream)
            if not set(fields).issubset(reader.fieldnames or []):
                raise ValueError('CSV anterior incompatible: faltan columnas. No se modificó.')
            for record in reader:
                if None in record or any(record.get(field) is None for field in fields):
                    continue
                try:
                    dua = normalize_dua(record['numero_del_dua'])
                except ValueError:
                    continue
                record = {field: record[field] for field in fields}
                record['numero_del_dua'] = dua
                imported += db.execute('''UPDATE jobs SET status='succeeded', result=?,
                    last_error='',error_kind='' WHERE id=? AND status!='succeeded' ''',
                    (json.dumps(record, ensure_ascii=False), dua)).rowcount
            self._set_meta(db, 'legacy_migrated', True)
        return imported

    def begin_run(self):
        with self.transaction() as db:
            db.execute("UPDATE attempts SET outcome='interrupted', finished=CURRENT_TIMESTAMP "
                       "WHERE outcome='running'")
            db.execute("UPDATE jobs SET status='pending', last_error='Ejecución interrumpida', "
                       "error_kind='interrupted' WHERE status='running'")
            db.execute('UPDATE jobs SET run_attempts=0')

    def claim(self, round_number):
        """Reserva atómica, una oportunidad por DUA en cada ronda."""
        with self.transaction() as db:
            row = db.execute('''SELECT id,dua FROM jobs INDEXED BY jobs_ready
                WHERE status IN ('pending','retry') AND run_attempts < ?
                ORDER BY run_attempts, rowid LIMIT 1''', (round_number,)).fetchone()
            if row is None:
                return None
            db.execute("UPDATE jobs SET status='running', attempts=attempts+1, "
                       "run_attempts=run_attempts+1 WHERE id=?", (row['id'],))
            attempt = db.execute('INSERT INTO attempts(job_id) VALUES (?)', (row['id'],)).lastrowid
            return {'id': row['id'], 'dua': row['dua'], 'attempt_id': attempt}

    def complete(self, job, record):
        if normalize_dua(record['numero_del_dua']) != job['dua']:
            raise ValueError('El resultado no corresponde al DUA reservado')
        serialized = json.dumps(record, ensure_ascii=False, allow_nan=False)
        with self.transaction() as db:
            self._check_claim(db, job)
            db.execute("UPDATE jobs SET status='succeeded',result=?,last_error='',"
                       "error_kind='',evidence='' WHERE id=?", (serialized, job['id']))
            db.execute("UPDATE attempts SET outcome='success',finished=CURRENT_TIMESTAMP WHERE id=?",
                       (job['attempt_id'],))

    @staticmethod
    def _check_claim(db, job):
        active = db.execute("SELECT 1 FROM attempts a JOIN jobs j ON a.job_id=j.id "
                            "WHERE a.id=? AND j.id=? AND a.outcome='running' AND j.status='running'",
                            (job['attempt_id'], job['id'])).fetchone()
        if not active:
            raise RuntimeError('Reserva vencida o ya resuelta')

    def retry(self, job, error, kind='extraction', evidence='', cancelled=False):
        with self.transaction() as db:
            self._check_claim(db, job)
            db.execute("UPDATE jobs SET status=?,last_error=?,error_kind=?,evidence=? WHERE id=?",
                       ('pending' if cancelled else 'retry', str(error)[:4000], kind, evidence, job['id']))
            db.execute('UPDATE attempts SET outcome=?,error=?,finished=CURRENT_TIMESTAMP WHERE id=?',
                       ('interrupted' if cancelled else 'retry', str(error)[:4000], job['attempt_id']))

    def summary(self):
        with self.lock:
            revisions = active_revisions(self.db)
            marks = ','.join('?' for _ in revisions)
            counts = dict(self.db.execute(f'''SELECT status,COUNT(*) FROM jobs
                WHERE status!='invalid' OR EXISTS (SELECT 1 FROM source_rows s
                    WHERE s.job_id=jobs.id AND s.revision IN ({marks})) GROUP BY status''',
                revisions).fetchall())
            attempts = self.db.execute('SELECT COALESCE(SUM(attempts),0) FROM jobs').fetchone()[0]
        return {**{state: counts.get(state, 0) for state in
                   ('pending', 'running', 'retry', 'succeeded', 'invalid')},
                'total': sum(counts.values()), 'attempts': attempts,
                'source_rows': self.meta('source_rows', 0)}

    def remove_jobs(self, ids=None, succeeded_only=False):
        """El llamador debe tener RunLock. Conserva resultados para futuras exportaciones."""
        if not succeeded_only and not ids:
            return 0
        with self.transaction() as db:
            db.execute('CREATE TEMP TABLE IF NOT EXISTS remove_ids (id TEXT PRIMARY KEY)')
            db.execute('DELETE FROM remove_ids')
            if succeeded_only:
                db.execute("INSERT INTO remove_ids SELECT id FROM jobs WHERE status='succeeded'")
            else:
                db.executemany('INSERT OR IGNORE INTO remove_ids VALUES (?)', ((key,) for key in ids))
            db.execute('''INSERT OR REPLACE INTO retained_results SELECT dua,result FROM jobs
                WHERE id IN (SELECT id FROM remove_ids) AND status='succeeded' ''')
            db.execute('DELETE FROM attempts WHERE job_id IN (SELECT id FROM remove_ids)')
            db.execute('DELETE FROM source_rows WHERE job_id IN (SELECT id FROM remove_ids)')
            removed = db.execute('DELETE FROM jobs WHERE id IN (SELECT id FROM remove_ids)').rowcount
            db.execute('DELETE FROM remove_ids')
            return removed

    def export(self, paths, fields):
        """Snapshots reemplazados atómicamente. Si falla un export, SQLite sigue intacto."""
        with self.lock:
            with atomic_output(paths['output_csv'], 'utf-8-sig', newline='') as stream:
                writer = csv.DictWriter(stream, fieldnames=fields)
                writer.writeheader()
                for row in self.db.execute('''SELECT result FROM retained_results WHERE dua NOT IN
                    (SELECT dua FROM jobs WHERE status='succeeded') UNION ALL
                    SELECT result FROM jobs WHERE status='succeeded' '''):
                    writer.writerow(json.loads(row[0]))
            with atomic_output(paths['processed_log']) as stream:
                for row in self.db.execute("SELECT dua FROM retained_results UNION SELECT dua FROM jobs WHERE status='succeeded'"):
                    stream.write(row[0] + '\n')
            pending = []
            revisions = active_revisions(self.db)
            marks = ','.join('?' for _ in revisions)
            for row in self.db.execute(f'''SELECT id,dua,status,attempts,last_error,error_kind,evidence
                    FROM jobs WHERE status!='succeeded' AND (status!='invalid' OR EXISTS (
                        SELECT 1 FROM source_rows s WHERE s.job_id=jobs.id AND s.revision IN ({marks})))
                    ORDER BY rowid''', revisions):
                item = dict(row)
                item['source_rows'] = [dict(source) for source in self.db.execute(
                    'SELECT revision,row_number,raw FROM source_rows WHERE job_id=?', (row['id'],))]
                pending.append(item)
            with atomic_output(paths['failed_duas']) as stream:
                stream.writelines(item['dua'] + '\n' for item in pending if item['dua'])
            with atomic_output(paths['pending_json']) as stream:
                json.dump({'summary': self.summary(), 'pending': pending}, stream,
                          ensure_ascii=False, indent=2)
